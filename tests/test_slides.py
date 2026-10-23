"""The English deck (D114): 15 slides from the facts, layout rules and reproducibility.

The facts come from the SYNTHETIC roots of test_report.py; no test reads data/.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("pptx")

from pptx import Presentation  # noqa: E402
from test_report import make_root  # noqa: E402

from traffic_twin.report.facts import collect_facts, load_facts  # noqa: E402
from traffic_twin.report.slides import (  # noqa: E402
    FOOTER,
    N_SLIDES,
    TITLE_PT,
    TITLE_WIDTH_IN,
    DeckError,
    a05_headline,
    build_deck,
    fits,
    title,
)

ROOT = Path(__file__).resolve().parents[1]
COMMITTED_DECK = ROOT / "docs" / "slides" / "urban_traffic_twin.pptx"
COMMITTED_FACTS = ROOT / "docs" / "report" / "facts.json"


@pytest.fixture
def pending_deck(tmp_path) -> Path:
    facts = collect_facts(make_root(tmp_path))
    return build_deck(facts, tmp_path / "deck.pptx", root=ROOT)


@pytest.fixture
def part8_deck(tmp_path) -> Path:
    facts = collect_facts(make_root(tmp_path, part8=True, sessions=3))
    return build_deck(facts, tmp_path / "deck8.pptx", root=ROOT)


def _texts(slide) -> list[str]:
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def _names(slide) -> set[str]:
    return {sh.name for sh in slide.shapes}


def test_fifteen_slides_inside_the_page_with_notes_and_footer(pending_deck):
    prs = Presentation(str(pending_deck))
    assert len(prs.slides) == N_SLIDES
    w, h = prs.slide_width, prs.slide_height
    assert abs(w / h - 16 / 9) < 0.01
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            assert sh.left >= 0 and sh.top >= 0, (i, sh.name)
            assert sh.left + sh.width <= w and sh.top + sh.height <= h, (i, sh.name)
        assert s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip(), i
        if i > 1:  # every content slide
            assert any(t == FOOTER for t in _texts(s)), i


def test_no_template_text_left(pending_deck):
    prs = Presentation(str(pending_deck))
    for i, s in enumerate(prs.slides, 1):
        for t in _texts(s) + [s.notes_slide.notes_text_frame.text]:
            assert not re.search(r"\{\{|\}\}|TODO|lorem|None\b", t), (i, t)


def test_titles_fit_on_one_line(pending_deck, part8_deck):
    for deck in (pending_deck, part8_deck):
        for s in Presentation(str(deck)).slides:
            for sh in s.shapes:
                if sh.name == "title":
                    assert fits(sh.text_frame.text), sh.text_frame.text


def test_a_title_that_would_wrap_is_a_build_error():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    with pytest.raises(DeckError, match=f"wraps at {TITLE_PT} pt"):
        title(slide, "A title far too long to fit on one line of the slide at this size " * 2)
    assert fits("Short title", TITLE_PT, TITLE_WIDTH_IN)


def test_s3_is_pending_without_the_part8_facts(pending_deck):
    prs = Presentation(str(pending_deck))
    results, rl = prs.slides[6], prs.slides[9]
    assert "s3-pending-card" in _names(results) and "a06-pending-card" in _names(rl)
    chart = next(sh for sh in results.shapes if sh.has_chart).chart
    assert "S3 RL" not in list(chart.plots[0].categories)
    assert not any("S3" in t and re.search(r"\d+\.\d\s?s", t) for t in _texts(results))


def test_s3_appears_with_the_part8_facts(part8_deck):
    prs = Presentation(str(part8_deck))
    results = prs.slides[6]
    chart = next(sh for sh in results.shapes if sh.has_chart).chart
    assert list(chart.plots[0].categories)[-1] == "S3 RL"
    labels = [p.data_label.text_frame.text for p in chart.plots[0].series[0].points]
    assert labels[-1] == "38.0\u00a0s"  # literal text, '.' decimal point
    assert "s3-pending-card" not in _names(results) and "a06-card" in _names(results)
    acc = prs.slides[12]
    assert "acc-status-A-10" in _names(acc)
    assert any(t == "DONE" for t in _texts(acc))


def _parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist() if not n.startswith("ppt/embeddings/")}


def test_rebuild_gives_the_same_parts_except_embedded_workbooks(tmp_path):
    facts = collect_facts(make_root(tmp_path))
    a = build_deck(facts, tmp_path / "a.pptx", root=ROOT)
    b = build_deck(facts, tmp_path / "b.pptx", root=ROOT)
    assert _parts(a) == _parts(b)


@pytest.mark.skipif(
    not (COMMITTED_DECK.is_file() and COMMITTED_FACTS.is_file()),
    reason="the deck is generated in integration (tasks.py slides)",
)
def test_committed_deck_matches_a_rebuild(tmp_path):
    rebuilt = build_deck(load_facts(COMMITTED_FACTS), tmp_path / "deck.pptx", root=ROOT)
    assert _parts(rebuilt) == _parts(COMMITTED_DECK)


def test_ambulance_headline_follows_the_sign_of_the_shift():
    assert "36\u00a0s sooner" in a05_headline(-36.5, "PASS")
    later = a05_headline(12.0, "FAIL")
    assert "12\u00a0s later" in later and "sooner" not in later
    assert "no change" in a05_headline(0.2, "FAIL")
    assert a05_headline(None, "PENDING") == "Ambulance (A-05): PENDING"
    for h in (a05_headline(-36.5, "PASS"), later):
        assert fits(h), h


def test_demo_slide_shows_the_presenter_rehearsal_apart_from_a08_a09(tmp_path):
    import yaml

    deck = Presentation(
        str(build_deck(collect_facts(make_root(tmp_path)), tmp_path / "a.pptx", root=ROOT))
    )
    text = " ".join(_texts(deck.slides[14]))
    assert "Automated: A-08 PENDING · A-09 PENDING" in text
    assert "Rehearsed by presenter: PENDING" in text
    root = make_root(tmp_path / "b")
    log = root / "docs/demo/rehearsal_log.yaml"
    log.parent.mkdir(parents=True, exist_ok=True)
    entry = {"date": "2026-11-02", "commit": "abc1234", "manual_code_edits": 0}
    log.write_text(yaml.safe_dump({"entries": [entry]}), encoding="utf-8")
    deck = Presentation(str(build_deck(collect_facts(root), tmp_path / "b.pptx", root=ROOT)))
    assert "Rehearsed by presenter: DONE" in " ".join(_texts(deck.slides[14]))
