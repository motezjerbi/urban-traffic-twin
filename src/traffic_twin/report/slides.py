"""docs/slides/urban_traffic_twin.pptx: the 15-slide English deck (D114, plan section 8).

Built with python-pptx from the facts (report.facts) only:

- 16:9 (13.33 x 7.5 in); bar comparisons are native charts whose data labels are written
  as literal text (no locale decimal comma); PNG only for the forest / CI figures and the
  app screenshots;
- every headline that states a result is computed from the facts; a fact that is not
  there yet (the Part 8 run, the user tests, the rehearsal) shows an amber PENDING tile,
  never a number;
- titles are measured (Calibri Bold, else DejaVu Sans Bold scaled) and a title that would
  wrap is a build error (DeckError);
- every content slide carries the footer 'Simulated data · ASSUMED signal program' and
  every slide has speaker notes;
- core properties use the meta.yaml date, so a rebuild gives the same XML parts except
  the charts' embedded workbooks (ppt/embeddings/*).
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LABEL_POSITION,
    XL_LEGEND_POSITION,
    XL_TICK_LABEL_POSITION,
)
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from traffic_twin.paths import APP_FIGURES, SLIDES_PPTX
from traffic_twin.report.acceptance import acceptance

# palette: asphalt (dominant), green-wave teal (support), signal amber (accent: PENDING),
# signal red only for a failed criterion
ASPHALT = RGBColor(0x17, 0x22, 0x2B)
ASPHALT_2 = RGBColor(0x24, 0x33, 0x3F)
TEAL = RGBColor(0x0B, 0x6E, 0x69)
TEAL_L = RGBColor(0x5F, 0xB8, 0xB0)
AMBER = RGBColor(0xF2, 0xA9, 0x00)
RED = RGBColor(0xC6, 0x3D, 0x2F)
INK = RGBColor(0x1F, 0x2A, 0x33)
MUTED = RGBColor(0x5B, 0x67, 0x70)
GREY = RGBColor(0xA9, 0xB6, 0xBE)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
TINT = RGBColor(0xEE, 0xF3, 0xF5)
GRID = RGBColor(0xE1, 0xE6, 0xE9)
FONT = "Calibri"
SERIES = {
    "S0": GREY,
    "S1": RGBColor(0x7C, 0x8C, 0x96),
    "S2": TEAL,
    "S3": TEAL_L,
    "S2_noP": RGBColor(0xC9, 0xD2, 0xD8),
    "S4": RGBColor(0xB4, 0xC8, 0xBE),
}
LABELS = {"S0": "S0 fixed", "S1": "S1 actuated", "S2": "S2 rules", "S3": "S3 RL"}

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.6)
TITLE_PT = 32
TITLE_WIDTH_IN = 12.1
FOOTER = "Simulated data · ASSUMED signal program"
N_SLIDES = 15


class DeckError(RuntimeError):
    """The deck breaks a layout rule (a title that wraps)."""


# -- facts access ----------------------------------------------------------------------
class F:
    def __init__(self, facts: dict, root: Path = Path(".")):
        self.facts = facts
        self.root = Path(root)

    def t(self, key: str) -> str | None:
        f = self.facts.get(key)
        return None if f is None else f.get("text")

    def v(self, key: str):
        f = self.facts.get(key)
        return None if f is None else f.get("value")

    def tp(self, key: str) -> str:
        """The text, or PENDING."""
        return self.t(key) or "PENDING"

    def part8(self) -> bool:
        return self.v("a06.status") in ("PASS", "INVALID")


def num(x: float, nd: int = 1, unit: str = "", signed: bool = False) -> str:
    s = f"{x:+.{nd}f}" if signed else f"{x:.{nd}f}"
    if float(s) == 0:
        s = s.lstrip("+-")
    return f"{s}\u00a0{unit}" if unit else s


# -- drawing helpers -------------------------------------------------------------------
_TITLE_FONTS = (
    ("C:/Windows/Fonts/calibrib.ttf", 1.0),
    ("/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf", 1.0),  # metric-compatible
)


def title_width_in(text: str, size_pt: float = TITLE_PT) -> float:
    """Rendered width of a bold title in inches (Calibri Bold or a metric-compatible
    font; else DejaVu Sans Bold scaled to Calibri's narrower design)."""
    from PIL import ImageFont

    for path, scale in _TITLE_FONTS:
        if Path(path).is_file():
            font = ImageFont.truetype(path, 1000)
            return font.getlength(text) / 1000 * size_pt / 72 * scale
    import matplotlib

    dejavu = Path(matplotlib.get_data_path()) / "fonts" / "ttf" / "DejaVuSans-Bold.ttf"
    font = ImageFont.truetype(str(dejavu), 1000)
    return font.getlength(text) / 1000 * size_pt / 72 * 0.86


def fits(text: str, size_pt: float = TITLE_PT, width_in: float = TITLE_WIDTH_IN) -> bool:
    return title_width_in(text, size_pt) <= width_in


def bg(slide, color) -> None:
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(
    slide,
    x,
    y,
    w,
    h,
    runs,
    *,
    size=16,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    name="text",
):
    """`runs`: a str, or a list of paragraphs, each a str or a list of (text, opts)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        items = [(para, {})] if isinstance(para, str) else para
        for t, o in items:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", False)
            r.font.color.rgb = o.get("color", color)
            if "space_after" in o:
                p.space_after = Pt(o["space_after"])
    return tb


def card(slide, x, y, w, h, fill, name, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.name = name
    s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    return s


def lamp(slide, x, y, d, color, name):
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    o.name = name
    o.fill.solid()
    o.fill.fore_color.rgb = color
    o.line.fill.background()
    o.shadow.inherit = False
    return o


def signal_head(slide, x, y, d, lit: int, name="signal"):
    """The deck's motif: a three-aspect signal head, one aspect lit (0 red, 1 amber,
    2 green)."""
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x - int(d * 0.25),
        y - int(d * 0.25),
        int(d * 1.5),
        int(d * 4.25),
    )
    box.name = f"{name}-box"
    box.fill.solid()
    box.fill.fore_color.rgb = ASPHALT_2
    box.line.fill.background()
    box.shadow.inherit = False
    for i, col in enumerate((RED, AMBER, TEAL_L)):
        lamp(slide, x, y + i * int(d * 1.3), d, col if i == lit else ASPHALT, f"{name}-{i}")


def status_color(status: str | None) -> RGBColor:
    s = (status or "PENDING").upper()
    if s in ("PASS", "DONE"):
        return TEAL_L
    if s in ("FAIL", "INVALID"):
        return RED
    return AMBER


def title(slide, t: str, sub: str | None = None, *, dark: bool = False) -> None:
    if not fits(t):
        raise DeckError(f"title wraps at {TITLE_PT} pt: {t!r}")
    text(
        slide,
        M,
        Inches(0.42),
        W - 2 * M,
        Inches(0.75),
        t,
        size=TITLE_PT,
        bold=True,
        color=PAPER if dark else INK,
        name="title",
    )
    if sub:
        text(
            slide,
            M,
            Inches(1.15),
            W - 2 * M,
            Inches(0.5),
            sub,
            size=16,
            color=TEAL_L if dark else MUTED,
            name="subtitle",
        )


def footer(slide, n: int, dark: bool = False) -> None:
    c = TEAL_L if dark else MUTED
    text(
        slide, M, H - Inches(0.42), Inches(9), Inches(0.3), FOOTER, size=11, color=c, name="footer"
    )
    text(
        slide,
        W - M - Inches(1),
        H - Inches(0.42),
        Inches(1),
        Inches(0.3),
        str(n),
        size=11,
        color=c,
        align=PP_ALIGN.RIGHT,
        name="page",
    )


def notes(slide, t: str) -> None:
    slide.notes_slide.notes_text_frame.text = t


def pending_tile(slide, x, y, w, h, head: str, body: str, name: str = "pending") -> None:
    card(slide, x, y, w, h, ASPHALT, f"{name}-card")
    lamp(slide, x + Inches(0.3), y + Inches(0.32), Inches(0.3), AMBER, f"{name}-lamp")
    text(
        slide,
        x + Inches(0.8),
        y + Inches(0.25),
        w - Inches(1.05),
        h - Inches(0.4),
        [
            [(head, {"size": 18, "bold": True, "color": PAPER, "space_after": 4})],
            [(body, {"size": 14, "color": TEAL_L})],
        ],
        name=f"{name}-text",
    )


def picture(slide, path: Path, x, y, max_w, max_h, name: str):
    from PIL import Image as PILImage

    with PILImage.open(path) as im:
        iw, ih = im.size
    sc = min(max_w / iw, max_h / ih)
    pic = slide.shapes.add_picture(str(path), x, y, int(iw * sc), int(ih * sc))
    pic.name = name
    return pic


def bar_chart(
    slide,
    x,
    y,
    w,
    h,
    categories,
    series: list[tuple[str, list[float], list[str]]],
    *,
    colors=None,
    name="chart",
    legend=False,
    per_point_colors=None,
    label_size=13,
):
    """A native clustered column chart; `series` = [(name, values, label texts)]."""
    cd = CategoryChartData()
    cd.categories = categories
    for sname, values, _ in series:
        cd.add_series(sname, values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, cd)
    gf.name = name
    ch = gf.chart
    ch.has_title = False  # the slide title and subtitle say what is plotted
    ch.font.name = FONT
    ch.font.size = Pt(13)
    ch.font.color.rgb = MUTED
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(13)
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.overlap = -10
    plot.has_data_labels = True
    for si, (_, _, labels) in enumerate(series):
        ser = plot.series[si]
        if colors:
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = colors[si]
        for i, lab in enumerate(labels):
            pt = ser.points[i]
            if per_point_colors:
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = per_point_colors[i]
            tf = pt.data_label.text_frame
            tf.text = lab  # literal text: no locale decimal comma
            r = tf.paragraphs[0].runs[0]
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = (
                Pt(label_size),
                True,
                INK,
                FONT,
            )
            pt.data_label.position = XL_LABEL_POSITION.OUTSIDE_END
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = GRID
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(12)
    ch.category_axis.tick_labels.font.size = Pt(13)
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW  # below negative bars
    ch.category_axis.format.line.color.rgb = RGBColor(0xC9, 0xD2, 0xD8)
    return gf


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# -- slides ----------------------------------------------------------------------------
def s01_title(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, ASPHALT)
    signal_head(s, M + Inches(0.1), Inches(1.6), Inches(0.42), lit=2)
    text(
        s,
        M + Inches(1.1),
        Inches(1.4),
        Inches(10.8),
        Inches(1.9),
        [
            [
                (
                    f.t("meta.title") or "Urban Traffic Twin",
                    {"size": 54, "bold": True, "color": PAPER},
                )
            ],
            [(f.t("meta.subtitle") or "", {"size": 22, "color": TEAL_L})],
        ],
        name="deck-title",
    )
    text(
        s,
        M + Inches(1.1),
        Inches(3.75),
        Inches(10.8),
        Inches(1.4),
        [
            [(f.t("meta.site") or "", {"size": 18, "color": PAPER, "space_after": 6})],
            [
                (
                    "Cars · buses · an ambulance · pedestrians: fixed, actuated, rules and RL "
                    "control "
                    "compared on the same seeds",
                    {"size": 16, "color": TEAL_L},
                )
            ],
        ],
        name="deck-sub",
    )
    who = f.t("meta.authors") or ""
    text(
        s,
        M + Inches(1.1),
        Inches(5.75),
        Inches(10.8),
        Inches(0.9),
        [
            [(who, {"size": 16, "color": PAPER})],
            [(f"{f.tp('meta.version')} · {f.tp('meta.date')}", {"size": 14, "color": TEAL_L})],
        ],
        name="authors",
    )
    text(
        s,
        M + Inches(1.1),
        H - Inches(0.6),
        Inches(11),
        Inches(0.3),
        f"{FOOTER} · field visit {f.tp('meta.field_visit')} not yet done",
        size=11,
        color=TEAL_L,
        name="footer",
    )
    notes(
        s,
        "Open with the four people of spec 1.1 stuck in the same morning peak. One real "
        "Paris junction, modelled and calibrated, four signal strategies compared fairly, "
        "and an app that tells each user what matters to them. Say early that every "
        "number is simulated and that the signal program is assumed until the field "
        "visit.",
    )


def s02_problem(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, PAPER)
    title(
        s,
        "Four users, one junction, one question",
        "Can adaptive signal control cut every user's delay without sacrificing one of them?",
    )
    users = [
        (
            "Student",
            "Will I make the exam? A travel time with an honest interval and a departure advice.",
        ),
        (
            "Ambulance",
            "Clear the junction now: preemption of the signals on the way to the hospital gate.",
        ),
        ("Bus driver", "Am I on schedule? Delay against the timetable, headway, priority."),
        (
            "Pedestrian",
            "Can I cross? A verdict for my walking speed, and a wait that stays bounded.",
        ),
    ]
    gap = Inches(0.3)
    cw = int((W - 2 * M - 3 * gap) / 4)
    for i, (who, what) in enumerate(users):
        x = M + i * (cw + gap)
        card(s, x, Inches(2.0), cw, Inches(3.3), TINT, f"user-{i}")
        lamp(
            s,
            x + Inches(0.3),
            Inches(2.3),
            Inches(0.3),
            (RED, AMBER, TEAL_L, TEAL)[i],
            f"user-lamp-{i}",
        )
        text(
            s,
            x + Inches(0.3),
            Inches(2.8),
            cw - Inches(0.6),
            Inches(0.5),
            who,
            size=22,
            bold=True,
            name=f"user-name-{i}",
        )
        text(
            s,
            x + Inches(0.3),
            Inches(3.4),
            cw - Inches(0.6),
            Inches(1.8),
            what,
            size=15,
            color=INK,
            name=f"user-need-{i}",
        )
    text(
        s,
        M,
        Inches(5.6),
        W - 2 * M,
        Inches(0.9),
        [
            [
                (
                    "The answer is measured per user, on the same demand and the same seeds for "
                    "every "
                    "controller, with pre-registered statistics (spec 13.3).",
                    {"size": 17, "color": TEAL},
                )
            ]
        ],
        name="method",
    )
    footer(s, n)
    notes(
        s,
        "These are the personas of the spec. The fifth profile, the operator, compares "
        "the strategies. The equity question is the hard part: a controller can lower the "
        "average delay while one user group pays for it, so every result is broken down "
        "per user and checked for hidden degradation.",
    )


def s03_site(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, PAPER)
    title(
        s,
        "The site: Bd de l'Hôpital × Bd Saint-Marcel",
        "Paris 13e, next to the Pitié-Salpêtrière hospital, sensor node 302",
    )
    # schematic junction (not to scale)
    cx, cy = Inches(3.6), Inches(4.3)
    road = Inches(1.0)
    for name, x, y, w, h in (
        ("road-ns", cx - road // 2, Inches(1.85), road, Inches(4.9)),
        ("road-w", M, cy - road // 2, cx - M, road),
    ):
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        r.name = name
        r.fill.solid()
        r.fill.fore_color.rgb = ASPHALT_2
        r.line.fill.background()
        r.shadow.inherit = False
    bus = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, cx + road // 2, cy - Inches(0.18), Inches(2.2), Inches(0.36)
    )
    bus.name = "bus-road"
    bus.fill.solid()
    bus.fill.fore_color.rgb = TEAL
    bus.line.fill.background()
    bus.shadow.inherit = False
    signal_head(s, cx + Inches(0.75), cy - Inches(1.7), Inches(0.2), lit=2, name="junction")
    for label, x, y in (
        ("N  Bd de l'Hôpital (Austerlitz)", cx + Inches(0.65), Inches(1.95)),
        ("S  Bd de l'Hôpital (Place d'Italie)", cx + Inches(0.65), Inches(6.2)),
        ("W  Bd Saint-Marcel", M, cy - Inches(0.95)),
        ("NB bus road · hospital gate", cx + Inches(0.6), cy + Inches(0.3)),
    ):
        text(s, x, y, Inches(3.3), Inches(0.4), label, size=13, color=INK, name="leg")
    x0 = Inches(7.3)
    facts = [
        (
            "Signalised junction",
            "three general legs, two bus-only legs, four signalised crossings, one traffic light",
        ),
        ("Buses", "lines 91, 24 and 57 with the Saint-Marcel - La Pitié stop"),
        (
            "Ambulance",
            "N to S towards the hospital main gate, one per run (seven in the ambulance probe)",
        ),
        (
            "Sensors",
            "City of Paris permanent loops: hourly flow and occupancy, one per general approach",
        ),
        ("Simulated window", "07:30-09:30, two hours of the school-term Tue-Thu peak"),
    ]
    paras = []
    for head, body in facts:
        paras.append([(head, {"size": 16, "bold": True, "color": TEAL})])
        paras.append([(body, {"size": 15, "color": INK, "space_after": 8})])
    text(s, x0, Inches(1.95), W - x0 - M, Inches(4.8), paras, name="site-facts")
    footer(s, n)
    notes(
        s,
        "Chosen with the spec's weighted grid: the only candidate with usable sensors on "
        "its approaches, a hospital, three bus lines and busy crossings. The drawing is a "
        "schematic, not to scale. OSM has no crossing on the N arm.",
    )


def s04_pipeline(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, PAPER)
    title(
        s,
        "From open sensor data to five user views",
        "Six layers, each with a simpler fallback that already works",
    )
    layers = [
        ("Data", "Paris permanent sensors, typical school-term days"),
        ("Simulation", "SUMO network from OSM, demand calibrated per arm"),
        ("Control", "S0 fixed · S1 actuated · S2 rules · S3 RL, one Safety Guard"),
        ("Prediction", "ETA quantiles with a conformal 80 % interval"),
        ("API", "FastAPI on precomputed runs, one envelope, degraded modes"),
        ("App", "five profiles, replay, operator dashboard"),
    ]
    gap = Inches(0.22)
    cw = int((W - 2 * M - 5 * gap) / 6)
    y = Inches(2.3)
    for i, (head, body) in enumerate(layers):
        x = M + i * (cw + gap)
        card(s, x, y, cw, Inches(2.9), ASPHALT if i % 2 == 0 else ASPHALT_2, f"layer-{i}")
        text(
            s,
            x + Inches(0.18),
            y + Inches(0.25),
            cw - Inches(0.36),
            Inches(0.5),
            head,
            size=19,
            bold=True,
            color=PAPER,
            name=f"layer-name-{i}",
        )
        text(
            s,
            x + Inches(0.18),
            y + Inches(0.95),
            cw - Inches(0.36),
            Inches(1.8),
            body,
            size=14,
            color=TEAL_L,
            name=f"layer-body-{i}",
        )
        if i < 5:
            a = s.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, x + cw, y + Inches(1.45), x + cw + gap, y + Inches(1.45)
            )
            a.name = f"arrow-{i}"
            a.line.color.rgb = AMBER
            a.line.width = Pt(2.5)
    text(
        s,
        M,
        Inches(5.6),
        W - 2 * M,
        Inches(1.0),
        [
            [
                (
                    "Everything the app shows comes from the simulation, through the API, labelled "
                    "'données simulées' with the calibration date (ENF-05).",
                    {"size": 16, "color": INK},
                )
            ]
        ],
        name="honesty",
    )
    footer(s, n)
    notes(
        s,
        "Data, simulation, control, prediction, API, app. Every layer has a simpler "
        "fallback: rules if RL fails, the reference table if the ETA model is missing, "
        "degraded answers if a run is missing. Nothing is a live city feed.",
    )


def s05_reality(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, PAPER)
    title(
        s,
        f"Twin vs reality: flows {f.tp('val.a02.verdict')}, timing not yet measured",
        "Spec 13.4 asks for three checks before any comparison; one is done",
    )
    tiles = [
        (
            "A-02 flows",
            f"{f.tp('val.a02.pass')}/{f.tp('val.a02.rows')}",
            f"inbound arm × slot rows within ±{f.tp('val.flow_threshold_pct')} of the sensors",
            status_color(f.t("val.a02.verdict")),
        ),
        (
            "GEH < 5",
            f"{f.tp('val.geh.n_below')}/{f.tp('val.geh.n_counts')}",
            "distinct counts",
            status_color(
                "PASS"
                if f.v("val.geh.n_below") is not None
                and f.v("val.geh.n_below") == f.v("val.geh.n_counts")
                else None
            ),
        ),
        (
            "Signal timing",
            f.tp("val.signal_not_measured"),
            "rows NOT_MEASURED until the field visit",
            AMBER,
        ),
        (
            "Queues",
            f.tp("val.queue_not_measured"),
            "rows NOT_MEASURED until the field visit",
            AMBER,
        ),
    ]
    gap = Inches(0.3)
    cw = int((W - 2 * M - 3 * gap) / 4)
    for i, (head, big, small, col) in enumerate(tiles):
        x = M + i * (cw + gap)
        card(s, x, Inches(1.95), cw, Inches(2.5), TINT, f"tile-{i}")
        lamp(s, x + Inches(0.25), Inches(2.25), Inches(0.26), col, f"tile-lamp-{i}")
        text(
            s,
            x + Inches(0.65),
            Inches(2.18),
            cw - Inches(0.8),
            Inches(0.4),
            head,
            size=16,
            bold=True,
            name=f"tile-head-{i}",
        )
        text(
            s,
            x + Inches(0.25),
            Inches(2.75),
            cw - Inches(0.5),
            Inches(0.9),
            big,
            size=40,
            bold=True,
            name=f"tile-big-{i}",
        )
        text(
            s,
            x + Inches(0.25),
            Inches(3.65),
            cw - Inches(0.5),
            Inches(0.75),
            small,
            size=13,
            color=MUTED,
            name=f"tile-small-{i}",
        )
    text(
        s,
        M,
        Inches(4.75),
        W - 2 * M,
        Inches(1.9),
        [
            [("What is ASSUMED", {"size": 17, "bold": True, "color": AMBER, "space_after": 4})],
            [
                (
                    "The signal program (cycle and greens), turn shares, lane use, bus headways "
                    "and "
                    "dwell, pedestrian flows. The inbound flow check is close to circular: the "
                    "demand "
                    "is built from the same sensors. The field visit "
                    f"({f.tp('meta.field_visit')}) measures the program and the queues; every "
                    "result "
                    "is provisional until then.",
                    {"size": 15, "color": INK},
                )
            ],
        ],
        name="assumed",
    )
    footer(s, n)
    notes(
        s,
        "Be explicit: A-02 passes, but it mostly shows that the network carries the "
        "demand in the right time slots. The independent checks are the signal timing and "
        "the queues, which wait for the field visit. If the measured program differs, "
        "Parts 3 to 8 are rerun and the report gets a version 2.",
    )


def s06_controllers(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, PAPER)
    title(
        s,
        "Four controllers behind one Safety Guard",
        "Same demand, same seeds; the guard filters every stage request of every controller",
    )
    ctrls = [
        ("S0", "Fixed time", "the ASSUMED provisional program"),
        ("S1", "Actuated", "SUMO-style gap logic: extend a green while vehicles arrive"),
        ("S2", "Rules", "ambulance preemption, bus priority, box clearing, pedestrian protection"),
        ("S3", "Reinforcement learning", "a DQN policy choosing stay or advance every few seconds"),
    ]
    gap = Inches(0.3)
    cw = int((W - 2 * M - 3 * gap) / 4)
    for i, (lab, name, body) in enumerate(ctrls):
        x = M + i * (cw + gap)
        card(s, x, Inches(1.95), cw, Inches(2.55), TINT, f"ctrl-{i}")
        text(
            s,
            x + Inches(0.25),
            Inches(2.15),
            cw - Inches(0.5),
            Inches(0.6),
            lab,
            size=30,
            bold=True,
            color=TEAL,
            name=f"ctrl-label-{i}",
        )
        text(
            s,
            x + Inches(0.25),
            Inches(2.8),
            cw - Inches(0.5),
            Inches(0.45),
            name,
            size=17,
            bold=True,
            name=f"ctrl-name-{i}",
        )
        text(
            s,
            x + Inches(0.25),
            Inches(3.3),
            cw - Inches(0.5),
            Inches(1.1),
            body,
            size=14,
            color=MUTED,
            name=f"ctrl-body-{i}",
        )
    card(s, M, Inches(4.75), W - 2 * M, Inches(1.85), ASPHALT, "guard")
    text(
        s,
        M + Inches(0.35),
        Inches(4.92),
        W - 2 * M - Inches(0.7),
        Inches(1.6),
        [
            [
                (
                    "Safety Guard (EF-11)",
                    {"size": 18, "bold": True, "color": PAPER, "space_after": 4},
                )
            ],
            [
                (
                    "Clearance never cut · min green · strict stage order · max green · "
                    "anti-starvation · and last, the pedestrian shield: no crossing waits past the "
                    f"{f.tp('eval.a04.bound')} bound. Every run audits the phases it showed.",
                    {"size": 15, "color": TEAL_L},
                )
            ],
        ],
        name="guard-text",
    )
    footer(s, n)
    notes(
        s,
        "S2 is the rules controller of Part 4; S3 is the RL policy selected in Part 7. "
        "The guard and the preemption keep priority over every controller, RL included. "
        "The shield is what makes S2's longer greens safe for pedestrians.",
    )


def s07_results(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, PAPER)
    order = ["S0", "S1", "S2"] + (["S3"] if f.t("final.all.person_delay_mean.s3") else [])
    vals = {}
    for lb in order:
        v = f.v(f"final.all.person_delay_mean.{lb.lower()}")
        if v is None:
            v = f.v(f"eval.person_delay.{lb.lower()}")
        if v is not None:
            vals[lb] = float(v)
    order = [lb for lb in order if lb in vals]
    seeds = f.tp("eval.seeds")
    if "S0" in vals and "S2" in vals:
        rel = 100 * (vals["S2"] - vals["S0"]) / vals["S0"]
        head = f"Rules cut the delay per person by {num(abs(rel), 0, '%')} against fixed time"
        if rel > 0:
            head = f"Rules raise the delay per person by {num(rel, 0, '%')} against fixed time"
    else:
        head = "Delay per person: results PENDING"
    title(
        s,
        head,
        f"Mean delay per person (all users, occupancy-weighted), seeds {seeds}, "
        "paired, same demand",
    )
    if order:
        bar_chart(
            s,
            M,
            Inches(1.85),
            Inches(7.6),
            Inches(4.95),
            [LABELS[k] for k in order],
            [
                (
                    "Person delay (s)",
                    [round(vals[k], 1) for k in order],
                    [num(vals[k], 1, "s") for k in order],
                )
            ],
            per_point_colors=[SERIES[k] for k in order],
            name="chart-person-delay",
        )
    x = M + Inches(8.0)
    cw = W - x - M
    card(s, x, Inches(1.85), cw, Inches(2.3), TINT, "stat-card")
    if "S0" in vals and "S2" in vals:
        rel = 100 * (vals["S2"] - vals["S0"]) / vals["S0"]
        text(
            s,
            x + Inches(0.3),
            Inches(1.98),
            cw - Inches(0.6),
            Inches(1.1),
            num(rel, 0, "%", signed=True),
            size=58,
            bold=True,
            color=TEAL,
            name="stat",
        )
        text(
            s,
            x + Inches(0.3),
            Inches(3.15),
            cw - Inches(0.6),
            Inches(0.9),
            f"S2 vs S0: {num(vals['S0'], 1, 's')} → {num(vals['S2'], 1, 's')} per person",
            size=16,
            name="stat-label",
        )
    if f.part8() and f.t("a06.primary_sentence"):
        card(s, x, Inches(4.45), cw, Inches(2.35), ASPHALT, "a06-card")
        lamp(
            s,
            x + Inches(0.3),
            Inches(4.77),
            Inches(0.3),
            status_color(f.t("a06.status")),
            "a06-lamp",
        )
        text(
            s,
            x + Inches(0.8),
            Inches(4.7),
            cw - Inches(1.05),
            Inches(2.0),
            [
                [
                    (
                        f"A-06 {f.tp('a06.status')}: S3 on seeds 1-10",
                        {"size": 18, "bold": True, "color": PAPER},
                    )
                ],
                [(f.tp("a06.primary_sentence"), {"size": 14, "color": TEAL_L})],
            ],
            name="a06-text",
        )
    else:
        pending_tile(
            s,
            x,
            Inches(4.45),
            cw,
            Inches(2.35),
            "S3 (RL) on seeds 1-10",
            "PENDING: the single Part 8 run has not been made. This tile and the S3 "
            "bar fill themselves from its outputs.",
            "s3-pending",
        )
    footer(s, n)
    notes(
        s,
        "Same demand and seeds for every controller. The person delay weights bus riders "
        "by an assumed occupancy, so it is the headline that reflects every user. S2 is "
        "the main solution; S3 is shown as computed by the pre-registered Part 8 run, "
        "whatever its direction.",
    )


EQUITY_KPIS = (
    ("car", "time_loss_mean", "Car time loss"),
    ("bus", "delay_mean", "Bus delay"),
    ("pedestrian", "ped_curb_wait_mean", "Pedestrian curb wait"),
    ("ambulance", "travel_time_total", "Ambulance travel time"),
)


def _rel_vs_s0(f: F, user: str, kpi: str, lb: str) -> float | None:
    rel = f.v(f"final.{user}.{kpi}.{lb.lower()}.rel")
    if rel is not None:
        return float(rel)
    if lb == "S3":
        return None  # S3 only from the Part 8 outputs
    a, b = f.v(f"eval.kpi.{user}.{kpi}.s0"), f.v(f"eval.kpi.{user}.{kpi}.{lb.lower()}")
    if a in (None, 0) or b is None:
        return None
    return 100 * (float(b) - float(a)) / float(a)


def s08_equity(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, PAPER)
    labels = ["S1", "S2"] + (["S3"] if f.part8() else [])
    series = []
    cats = [c[2] for c in EQUITY_KPIS]
    for lb in labels:
        vals = [_rel_vs_s0(f, u, k, lb) for u, k, _ in EQUITY_KPIS]
        if any(v is None for v in vals):
            continue
        series.append(
            (LABELS[lb], [round(v, 1) for v in vals], [num(v, 1, "%", signed=True) for v in vals])
        )
    worst = None
    for name, vals, _ in series:
        for c, v in zip(cats, vals, strict=True):
            if v > 0 and (worst is None or v > worst[2]):
                worst = (name, c, v)
    if worst:
        change = num(worst[2], 1, "%", signed=True)
        head = f"Who pays? {worst[1]} worsens most: {change} ({worst[0]})"
        if not fits(head):
            head = f"Who pays? {worst[1]}: {change} vs S0"
    else:
        head = (
            "Who pays? No headline KPI worse than fixed time"
            if series
            else "Equity per user: PENDING"
        )
    title(s, head, "Change against S0 per user's headline KPI, seeds 1-10 (negative is better)")
    if series:
        bar_chart(
            s,
            M,
            Inches(1.85),
            Inches(8.3),
            Inches(4.95),
            cats,
            series,
            colors=[SERIES[lb] for lb in labels[: len(series)]],
            legend=True,
            name="chart-equity",
            label_size=11 if len(series) > 2 else 13,
        )
    x = M + Inches(8.7)
    cw = W - x - M
    text(
        s,
        x,
        Inches(1.95),
        cw,
        Inches(2.6),
        [
            [
                (
                    "Hidden degradation (13.3)",
                    {"size": 17, "bold": True, "color": INK, "space_after": 4},
                )
            ],
            [
                (
                    f"Flagged: {f.tp('eval.hidden_degradation')}. The person delay improves while "
                    "a "
                    "car arm, a bus flow or a crossing is more than "
                    f"{f.tp('eval.equity_pct')} worse than S0.",
                    {"size": 14, "color": INK},
                )
            ],
        ],
        name="hidden",
    )
    if not f.part8():
        pending_tile(
            s,
            x,
            Inches(4.7),
            cw,
            Inches(2.1),
            "S3 equity",
            "PENDING until the Part 8 run (tables vs S0, S1, S2).",
            "s3-equity",
        )
    footer(s, n)
    notes(
        s,
        "Equity is checked on every KPI against S0, per car arm, bus flow and crossing. "
        "The pedestrian cost is the recurring trade-off: longer greens for vehicles mean "
        "longer curb waits, kept under the bound by the shield.",
    )


def a05_headline(hl, verdict: str) -> str:
    """The A-05 headline; 'sooner' or 'later' follows the sign of the HL shift (S2 - S0)."""
    if hl is None:
        return "Ambulance (A-05): PENDING"
    hl = float(hl)
    if round(hl) == 0:
        return f"Preemption: no change in ambulance travel time (A-05 {verdict})"
    when = "sooner" if hl < 0 else "later"
    return f"Preemption: the ambulance arrives {num(abs(hl), 0, 's')} {when} (A-05 {verdict})"


def s09_ambulance(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, PAPER)
    head = a05_headline(f.v("eval.a05.hl"), f.tp("eval.a05.verdict"))
    title(
        s,
        head,
        "Travel time from the planned departure, main family, one ambulance per run, seeds 1-10",
    )
    cats, vals, cols = [], [], []
    for lb, key in (
        ("S0", "eval.a05.mean_s0"),
        ("S2", "eval.a05.mean_s2"),
        ("S3", "final.ambulance.travel_time_total.s3"),
    ):
        if f.v(key) is not None:
            cats.append(LABELS[lb])
            vals.append(float(f.v(key)))
            cols.append(SERIES[lb])
    if cats:
        bar_chart(
            s,
            M,
            Inches(1.85),
            Inches(6.4),
            Inches(4.95),
            cats,
            [("Travel time (s)", [round(v, 1) for v in vals], [num(v, 1, "s") for v in vals])],
            per_point_colors=cols,
            name="chart-ambulance",
        )
    x = M + Inches(6.9)
    cw = W - x - M
    rows = [
        (
            "Main family, S2 vs S0 (A-05)",
            f"HL {f.tp('eval.a05.hl')} {f.tp('eval.a05.ci')}, one-sided p {f.tp('eval.a05.p')}",
        ),
        (
            "Probe (7 ambulances), S2 vs S0",
            f"HL {f.tp('eval.a05.probe.hl')}, p {f.tp('eval.a05.probe.p')}",
        ),
        (
            "Probe, S2 vs S2 without preemption",
            f"HL {f.tp('eval.a05.probe_nop.hl')}, p {f.tp('eval.a05.probe_nop.p')}",
        ),
    ]
    paras = []
    for head2, body in rows:
        paras.append([(head2, {"size": 16, "bold": True, "color": TEAL})])
        paras.append([(body, {"size": 15, "color": INK, "space_after": 10})])
    paras.append(
        [
            (
                "One ambulance per run gives the main test little power; the probe family "
                "attributes the gain to preemption.",
                {"size": 14, "color": MUTED},
            )
        ]
    )
    text(s, x, Inches(1.95), cw, Inches(4.8), paras, name="a05-detail")
    footer(s, n)
    notes(
        s,
        "A-05 is the pre-registered one-sided test on the main family. The probe family "
        "runs seven ambulances per run and compares S2 with and without preemption, which "
        "isolates the effect. The ambulance obeys red lights in the model, so preemption "
        "is the only way through.",
    )


def s10_rl(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, PAPER)
    verdict = f.tp("rl7.verdict")
    title(
        s,
        f"RL learns, {'and' if verdict == 'PASS' else 'but'} the interim spec 10.6 verdict "
        f"is {verdict}",
        f"Part 7 dev-test seeds {f.tp('rl7.seeds')}, looked at once; S2 stays the main solution",
    )
    crit = [
        (
            "C1",
            "Person delay vs S1",
            f.tp("rl7.c1.hl"),
            f"HL, p {f.tp('rl7.c1.p')}",
            f.t("rl7.c1.pass"),
        ),
        (
            "C2",
            f"Equity vs S1 (≤ {f.tp('rl7.c2.limit')})",
            f.tp("rl7.c2.ped_rel"),
            "pedestrian curb wait",
            f.t("rl7.c2.pass"),
        ),
        ("C3", "Safety", f.tp("rl7.c3.violations"), "violations", f.t("rl7.c3.pass")),
        ("C4", "Reproducibility", f.tp("rl7.c4.text"), "picks below S1", f.t("rl7.c4.pass")),
    ]
    gap = Inches(0.25)
    cw = int((Inches(8.2) - 3 * gap) / 4)
    for i, (cid, label, big, small, ok) in enumerate(crit):
        x = M + i * (cw + gap)
        y = Inches(1.9)
        card(s, x, y, cw, Inches(2.35), TINT, f"crit-{cid}")
        col = status_color("PASS" if ok == "pass" else "FAIL" if ok == "fail" else None)
        lamp(s, x + Inches(0.2), y + Inches(0.27), Inches(0.24), col, f"lamp-{cid}")
        text(
            s,
            x + Inches(0.55),
            y + Inches(0.2),
            cw - Inches(0.7),
            Inches(0.75),
            [
                [(f"{cid} · {ok or 'pending'}", {"size": 14, "bold": True})],
                [(label, {"size": 12, "color": INK})],
            ],
            name=f"crit-head-{cid}",
        )
        text(
            s,
            x + Inches(0.2),
            y + Inches(1.1),
            cw - Inches(0.4),
            Inches(0.7),
            big,
            size=26,
            bold=True,
            name=f"crit-big-{cid}",
        )
        text(
            s,
            x + Inches(0.2),
            y + Inches(1.8),
            cw - Inches(0.4),
            Inches(0.45),
            small,
            size=12,
            color=MUTED,
            name=f"crit-small-{cid}",
        )
    if f.part8():
        # the test of record in large text (the Part 7 forest raster is unreadable at
        # slide size; its figure stays in the report and on the Part 7 page)
        lines = [
            [
                (
                    "A-06 on seeds 1-10, S3 − reference, person-weighted delay",
                    {"size": 15, "bold": True, "color": INK, "space_after": 4},
                )
            ]
        ]
        for ref, name in (("s1", "S1 actuated (primary)"), ("s2", "S2 rules"), ("s0", "S0 fixed")):
            o = f.tp(f"a06.{ref}.outcome")
            o = f"S3 {o}" if o in ("better", "worse") else o
            lines.append(
                [
                    (f"vs {name}: ", {"size": 15, "bold": True, "color": INK}),
                    (
                        f"HL {f.tp(f'a06.{ref}.hl')}, 95 % CI {f.tp(f'a06.{ref}.ci')}, "
                        f"Holm p {f.tp(f'a06.{ref}.p_holm')}: {o}",
                        {"size": 15, "color": INK},
                    ),
                ]
            )
        text(s, M, Inches(4.55), Inches(8.2), Inches(2.2), lines, name="a06-rows")
    else:
        fig = f.root / "docs/figures/rl7/s3_forest.png"
        if fig.is_file():
            picture(s, fig, M, Inches(4.4), Inches(8.2), Inches(2.4), "fig-s3-forest")
    x = M + Inches(8.55)
    cw = W - x - M
    if f.part8():
        body = [
            [
                (
                    f"A-06 on seeds 1-10: {f.tp('a06.status')}",
                    {"size": 17, "bold": True, "color": PAPER},
                )
            ],
            [(f.tp("a06.primary_sentence"), {"size": 14, "color": TEAL_L, "space_after": 6})],
            [
                (
                    f"Spec 10.6 check: seeds 1-10 {f.tp('rep106.confirms')} the interim verdict.",
                    {"size": 14, "color": PAPER},
                )
            ],
        ]
        card(s, x, Inches(1.9), cw, Inches(2.6), ASPHALT, "a06-card")
        text(s, x + Inches(0.3), Inches(2.1), cw - Inches(0.6), Inches(2.3), body, name="a06-text")
    else:
        pending_tile(
            s,
            x,
            Inches(1.9),
            cw,
            Inches(2.6),
            "A-06 on seeds 1-10",
            "PENDING: the test of record runs once, in Part 8, and is shown whatever "
            "its direction.",
            "a06-pending",
        )
    text(
        s,
        x,
        Inches(4.75),
        cw,
        Inches(2.0),
        [
            [("Why the fallback", {"size": 17, "bold": True, "color": INK, "space_after": 4})],
            [
                (
                    "S3 does not observe pedestrians; the shield alone keeps the bound, so curb "
                    "waits "
                    "drift up while cars and buses gain.",
                    {"size": 14, "color": INK},
                )
            ],
        ],
        name="why",
    )
    footer(s, n)
    notes(
        s,
        "The criteria were pre-registered (D79) and looked at once. C2 fails on the "
        "pedestrian curb wait, so the week-7 rule makes the rules controller the main "
        "solution and RL a study. A-06 is still run on seeds 1-10 as registered and "
        "presented whatever its direction; it does not reopen the choice.",
    )


def s11_eta(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, PAPER)
    title(
        s,
        f"An honest travel time: the 80 % interval covers {f.tp('eta.a07.coverage_pct')}",
        f"A-07 {f.tp('eta.a07.verdict')}: target band {f.tp('eta.a07.band')}, realistic test "
        "worlds generated after the model was frozen",
    )
    card(s, M, Inches(1.95), Inches(4.0), Inches(4.8), TINT, "eta-card")
    paras = [
        [("Coverage", {"size": 16, "bold": True, "color": TEAL})],
        [
            (
                f"{f.tp('eta.a07.coverage_pct')} {f.tp('eta.a07.ci')}",
                {"size": 22, "bold": True, "space_after": 10},
            )
        ],
        [("Median error (MAE)", {"size": 16, "bold": True, "color": TEAL})],
        [
            (
                f"{f.tp('eta.a07.mae')} vs {f.tp('eta.a07.ref_mae')} for the reference table",
                {"size": 18, "bold": True, "space_after": 10},
            )
        ],
        [("Trips", {"size": 16, "bold": True, "color": TEAL})],
        [(f"{f.tp('eta.a07.n')} trips in {f.tp('eta.a07.n_worlds')} test worlds", {"size": 16})],
    ]
    text(s, M + Inches(0.3), Inches(2.2), Inches(3.4), Inches(4.4), paras, name="eta-stats")
    fig = f.root / "docs/figures/eta/coverage_by_group.png"
    if fig.is_file():
        picture(
            s,
            fig,
            M + Inches(4.3),
            Inches(1.95),
            W - 2 * M - Inches(4.3),
            Inches(4.8),
            "fig-eta-coverage",
        )
    footer(s, n)
    notes(
        s,
        "XGBoost quantiles plus a conformal step calibrated on separate worlds. The "
        "student view always shows the interval and a departure advice. The model learns "
        "the simulator, so it is retrained after the field visit.",
    )


SHOTS = (
    ("home", "Home: five profiles"),
    ("pedestrian", "Pedestrian"),
    ("ambulance", "Ambulance (kiosk)"),
    ("replay", "Replay"),
)


def s12_app(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, PAPER)
    title(
        s,
        "One API, two front ends, every state handled",
        "User app (five profiles and replay) and the Streamlit operator view, both through the API",
    )
    gap = Inches(0.25)
    cw = int((Inches(8.4) - gap) / 2)
    ch = Inches(2.3)
    for i, (shot, label) in enumerate(SHOTS):
        x = M + (i % 2) * (cw + gap)
        y = Inches(1.85) + (i // 2) * (ch + Inches(0.15))
        path = f.root / APP_FIGURES / f"{shot}.png"
        if path.is_file():
            picture(s, path, x, y, cw, ch - Inches(0.3), f"shot-{shot}")
            text(
                s,
                x,
                y + ch - Inches(0.28),
                cw,
                Inches(0.28),
                label,
                size=12,
                color=MUTED,
                name=f"shot-label-{shot}",
            )
        else:
            pending_tile(
                s,
                x,
                y,
                cw,
                ch - Inches(0.1),
                label,
                "PENDING: screenshot taken by tasks.py app-shots",
                f"shot-{shot}",
            )
    x = M + Inches(8.75)
    cw2 = W - x - M
    items = [
        ("States (spec 12.1)", "loading · ok · degraded · stale · unavailable · error"),
        ("Honesty", "'données simulées' with the calibration date in every view"),
        ("At most 2 clicks", "from the home screen to every answer (ENF-06)"),
        ("A-08", f.tp("acceptance.A-08")),
    ]
    paras = []
    for head, body in items:
        paras.append([(head, {"size": 16, "bold": True, "color": TEAL})])
        paras.append([(body, {"size": 15, "space_after": 8})])
    text(s, x, Inches(1.95), cw2, Inches(4.8), paras, name="states")
    footer(s, n)
    notes(
        s,
        "The user app is Motez's design, fed only by the API; the operator view is "
        "Streamlit. No value is invented on the client: a missing answer shows a degraded "
        "or unavailable state with the reason.",
    )


def s13_acceptance(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, ASPHALT)
    rows = acceptance(f.facts)
    done = sum(1 for r in rows if r["status"] in ("PASS", "DONE"))
    title(
        s,
        f"Acceptance: {done} of {len(rows)} criteria met so far",
        "Spec 17, each status generated from its evidence file",
        dark=True,
    )
    gap = Inches(0.22)
    cw = int((W - 2 * M - 4 * gap) / 5)
    chh = Inches(2.3)
    for i, r in enumerate(rows):
        x = M + (i % 5) * (cw + gap)
        y = Inches(1.9) + (i // 5) * (chh + Inches(0.2))
        card(s, x, y, cw, chh, ASPHALT_2, f"acc-{r['id']}")
        lamp(
            s,
            x + Inches(0.22),
            y + Inches(0.25),
            Inches(0.36),
            status_color(r["status"]),
            f"acc-lamp-{r['id']}",
        )
        text(
            s,
            x + Inches(0.75),
            y + Inches(0.2),
            cw - Inches(0.9),
            Inches(0.5),
            [[(r["id"], {"size": 18, "bold": True, "color": PAPER})]],
            name=f"acc-id-{r['id']}",
        )
        text(
            s,
            x + Inches(0.22),
            y + Inches(0.75),
            cw - Inches(0.44),
            Inches(0.4),
            r["status"],
            size=16,
            bold=True,
            color=status_color(r["status"]),
            name=f"acc-status-{r['id']}",
        )
        text(
            s,
            x + Inches(0.22),
            y + Inches(1.2),
            cw - Inches(0.44),
            chh - Inches(1.3),
            SHORT[r["id"]],
            size=13,
            color=TEAL_L,
            name=f"acc-what-{r['id']}",
        )
    footer(s, n, dark=True)
    notes(s, acceptance_notes(rows, f.v("demo.human_rehearsals")))


def acceptance_notes(rows: list[dict], human_rehearsals) -> str:
    """Slide 13's speaker notes, generated from the same rows as its lamps: what is still
    PENDING is named from the evidence, never from a fixed list."""
    pending = [r for r in rows if r["status"] == "PENDING"]
    if pending:
        waits = "; ".join(
            f"{r['id']} ({SHORT.get(r['id'], r['id'])}): {r['evidence']}" for r in pending
        )
        say = f"PENDING means the evidence does not exist yet. Still pending: {waits}."
    else:
        say = "No criterion is PENDING."
    if not human_rehearsals:
        say += " The presenter's own rehearsal of the demo is not recorded yet."
    return (
        f"Each lamp is generated from its evidence file. {say} A-06 passes when the "
        "comparison is valid and shown, whatever its direction; A-10 needs three to five "
        "real users."
    )


SHORT = {
    "A-01": "same seed, same KPIs",
    "A-02": "simulated flow vs sensors",
    "A-03": "four controllers, one KPI file",
    "A-04": "no dangerous sequence, bounded pedestrian wait",
    "A-05": "preemption speeds up the ambulance",
    "A-06": "RL vs baselines, shown either way",
    "A-07": "ETA interval coverage",
    "A-08": "every view via the API, all states",
    "A-09": "demo without touching the code",
    "A-10": "3 to 5 real users",
}


def s14_limits(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, PAPER)
    title(s, "Limitations and what comes next")
    left = [
        "The twin learns the simulator: one junction, one day type, ten seeds.",
        "The signal program is ASSUMED until the field visit; timing and queues are NOT_MEASURED.",
        "Detection is ideal (pedestrians, buses, ambulance before it enters).",
        "S3 does not observe pedestrians; S4 (S3 without preemption) is out of distribution.",
        "One ambulance per run gives A-05 little power.",
    ]
    right = [
        f"Field visit {f.tp('meta.field_visit')}: measured program, queues, counts.",
        "If the program differs: rerun Parts 3 to 8, report version 2.",
        "Real user sessions (A-10), then fixes and a second round.",
        "RL that observes pedestrians, only as a new pre-registered study.",
    ]
    for i, (head, items, col) in enumerate((("Limitations", left, RED), ("Next", right, TEAL))):
        x = M + i * Inches(6.2)
        paras = [[(head, {"size": 20, "bold": True, "color": col, "space_after": 6})]]
        paras += [[("•  " + it, {"size": 16, "space_after": 8})] for it in items]
        text(s, x, Inches(1.6), Inches(5.9), Inches(5.2), paras, name=f"col-{i}")
    footer(s, n)
    notes(
        s,
        "Be direct about the limits: everything is simulated, the program is assumed, "
        "and the RL policy is blind to pedestrians. The field visit can change the "
        "numbers; the code and the pre-registered procedure stay.",
    )


def s15_demo(prs, f: F, n: int) -> None:
    s = _blank(prs)
    bg(s, ASPHALT)
    title(
        s,
        "Live demo",
        "tasks.py demo: the API and the operator view, started by one command",
        dark=True,
    )
    steps = [
        ("1", "Home", "five profiles, two clicks to any answer"),
        ("2", "Pedestrian", "crossing S at normal walking speed"),
        ("3", "Student", "Austerlitz to the hospital, exam at 08:18"),
        ("4", "Bus", "line 91 delay and headway"),
        ("5", "Ambulance", "kiosk screen, three junctions"),
        ("6", "Replay", "the S3 run on demo seed 4000"),
        ("7", "Operator", "comparison, tests, equity, curves"),
    ]
    paras = []
    for num_, head, body in steps:
        paras.append(
            [
                (f"{num_}  ", {"size": 17, "bold": True, "color": AMBER}),
                (f"{head}  ", {"size": 17, "bold": True, "color": PAPER}),
                (body, {"size": 16, "color": TEAL_L}),
            ]
        )
    text(s, M, Inches(1.85), Inches(7.6), Inches(2.6), paras, name="steps")
    x = M + Inches(8.0)
    cw = W - x - M
    card(s, x, Inches(1.85), cw, Inches(2.6), ASPHALT_2, "fallback")
    text(
        s,
        x + Inches(0.3),
        Inches(2.05),
        cw - Inches(0.6),
        Inches(2.3),
        [
            [("Fallback", {"size": 18, "bold": True, "color": AMBER, "space_after": 4})],
            [
                (
                    "If anything fails live: the screenshots of slides 12 and 15 and the "
                    "precomputed "
                    "replay. No simulation runs during the demo.",
                    {"size": 15, "color": PAPER},
                )
            ],
        ],
        name="fallback-text",
    )
    for i, (shot, label) in enumerate(
        (("replay", "Replay (fallback)"), ("streamlit", "Operator (fallback)"))
    ):
        sx = M + i * Inches(3.85)
        sy, sw, sh = Inches(4.55), Inches(3.6), Inches(2.25)
        path = f.root / APP_FIGURES / f"{shot}.png"
        if path.is_file():
            picture(s, path, sx, sy, sw, sh - Inches(0.3), f"fallback-{shot}")
            text(
                s,
                sx,
                sy + sh - Inches(0.28),
                sw,
                Inches(0.28),
                label,
                size=12,
                color=TEAL_L,
                name=f"fallback-label-{shot}",
            )
        else:
            pending_tile(
                s,
                sx,
                sy,
                sw,
                sh - Inches(0.1),
                label,
                "PENDING: screenshot taken by tasks.py app-shots",
                f"fallback-{shot}",
            )
    card(s, x, Inches(4.7), cw, Inches(2.05), ASPHALT_2, "rehearsal")
    text(
        s,
        x + Inches(0.3),
        Inches(4.9),
        cw - Inches(0.6),
        Inches(1.8),
        [
            [("Rehearsal", {"size": 18, "bold": True, "color": PAPER, "space_after": 4})],
            [
                (
                    f"Automated: A-08 {f.tp('demo.a08')} · A-09 {f.tp('demo.a09')}",
                    {"size": 15, "color": TEAL_L},
                )
            ],
            [
                (
                    f"Rehearsed by presenter: {f.tp('demo.rehearsed')}",
                    {"size": 15, "color": TEAL_L},
                )
            ],
        ],
        name="rehearsal-text",
    )
    footer(s, n, dark=True)
    notes(
        s,
        "Follow docs/demo/runbook.md. Say that the data are simulated, that the signal "
        "program is assumed, that RL is a study after the week-7 FALLBACK, and what A-06 "
        "showed. Switch to the fallback screenshots at the first sign of trouble.",
    )


SLIDES = (
    s01_title,
    s02_problem,
    s03_site,
    s04_pipeline,
    s05_reality,
    s06_controllers,
    s07_results,
    s08_equity,
    s09_ambulance,
    s10_rl,
    s11_eta,
    s12_app,
    s13_acceptance,
    s14_limits,
    s15_demo,
)


def _meta_date(f: F) -> dt.datetime:
    try:
        d = dt.date.fromisoformat(str(f.v("meta.date"))[:10])
    except ValueError:
        d = dt.date(2026, 1, 1)
    return dt.datetime(d.year, d.month, d.day, 12, 0, 0)


def build_deck(facts: dict, out: Path = SLIDES_PPTX, *, root: Path = Path(".")) -> Path:
    """Write the deck (15 slides) built from `facts` (figures read under `root`, the repo
    root); returns its path."""
    f = F(facts, root)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    for i, make in enumerate(SLIDES, 1):
        make(prs, f, i)
    assert len(prs.slides) == N_SLIDES
    cp = prs.core_properties
    when = _meta_date(f)
    cp.title = f"{f.t('meta.title') or 'Urban Traffic Twin'}: final presentation"
    cp.author = f.t("meta.authors") or "Urban Traffic Twin"
    cp.subject = "Digital twin of the Pitié-Salpêtrière junction (simulated data)"
    cp.keywords = "SUMO; traffic signal control; reinforcement learning; digital twin"
    cp.comments = f"Built from docs/report/facts.json, code {f.t('meta.code') or 'unknown'}"
    cp.last_modified_by = "traffic_twin.report.slides"
    cp.revision = 1
    cp.created = cp.modified = cp.last_printed = when
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out
